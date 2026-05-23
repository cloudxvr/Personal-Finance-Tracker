from BillingSystem import Ui_MainWindow
from Bill import Ui_Bill
from Statistic import Ui_Statistic
from Login import Ui_Login
from openpyxl import load_workbook, Workbook
import matplotlib.pyplot as plt
import time
import datetime
import random
from PyQt5.QtCore import Qt, QDate
from PyQt5 import QtWidgets
from PyQt5.QtWidgets import QMainWindow, QWidget, QCheckBox, QTableWidgetItem, QMessageBox, QLineEdit, QApplication
from PyQt5.QtGui import QPixmap
import os
import sys


class Login_window(QWidget, Ui_Login):
    user_file = ""

    def __init__(self):
        super(Login_window, self).__init__()
        self.setupUi(self)
        self.lineEdit_2.setEchoMode(QLineEdit.Password)
        self.pushButton.clicked.connect(self.login)
        self.pushButton_2.clicked.connect(self.register)

    def login(self):
        user = self.lineEdit.text()
        password = self.lineEdit_2.text()
        Users = load_workbook('Users.xlsx')
        sheet = Users.active
        flag = 0
        for i in range(1, sheet.max_row+1):
            if str(sheet.cell(i, 1).value) == user:
                flag = 1
                if str(sheet.cell(i, 2).value) == password:
                    self.user_file = user + ".xlsx"
                    self.hide()
                    self.main = myMainWindow()
                    self.main.show()
                else:
                    self.show_message("  密码错误    ")
                break
            else:
                flag = 0
        if flag == 0:
            self.show_message("  请先注册    ")

    def register(self):
        user = self.lineEdit.text()
        password = self.lineEdit_2.text()
        Users = load_workbook('Users.xlsx')
        sheet = Users.active
        flag = 1
        for i in range(1, sheet.max_row+1):
            if str(sheet.cell(i, 1).value) == user:
                self.show_message("  请勿重复注册    ")
                flag = 0
                break
        if flag:
            self.user_file = user + ".xlsx"
            sheet.append([user, password])
            Users.save('Users.xlsx')
            wb = Workbook()
            wb.save(self.user_file)
            self.hide()
            self.main = myMainWindow()
            self.main.show()

    def show_message(self, message):
        QMessageBox.information(self, "提示", message)


class myMainWindow(QMainWindow, Ui_MainWindow):
    items = ["餐饮", "购物", "日常", "交通", "学习", "医疗", "娱乐", "水电", "其他"]

    def __init__(self):
        super(myMainWindow, self).__init__()
        self.setupUi(self)
        self.pushButton.clicked.connect(self.getData)
        self.actionShow.triggered.connect(self.showBill)
        self.actionSearch.triggered.connect(self.showStatistic)
        self.actionUndo.triggered.connect(self.undo)
        self.pushButton_2.clicked.connect(self.importData)
        self.action1.triggered.connect(self.switch)
        self.comboBox.addItems(self.items)
        self.dateEdit.setDate(QDate.currentDate())

    def getData(self):
        money = self.lineEdit.text()
        type = self.comboBox.currentText()
        date = self.dateEdit.date().toString(Qt.ISODate)
        remark = self.lineEdit_3.text()
        AccountBook = load_workbook(login.user_file)
        sheet = AccountBook.active
        sheet.append([date, type, money, remark])
        AccountBook.save(login.user_file)
        self.lineEdit.clear()
        self.lineEdit_3.clear()

    def importData(self):
        wb = Workbook()
        sheet = wb.active
        for i in range(5):
            sheet.append([datetime.datetime.now().strftime("%Y-%m-%d"), random.choice(self.items),
                         str(random.randint(10, 500)), "无"])
        wb.save("ImportData.xlsx")
        path = self.filepath()
        if path:
            imp_wb = load_workbook(path)
            acc_wb = load_workbook(login.user_file)
            imp_sheet = imp_wb.active
            acc_sheet = acc_wb.active
            for i in range(1, imp_sheet.max_row+1):
                acc_sheet.append([imp_sheet.cell(i, 1).value, imp_sheet.cell(i, 2).value,
                                  imp_sheet.cell(i, 3).value, imp_sheet.cell(i, 4).value])
            acc_wb.save(login.user_file)

    def filepath(self):
        path, filetype = QtWidgets.QFileDialog.getOpenFileName(self, "选取文件", os.getcwd(),
                                                               "All Files(*);;Text Files(*.txt)")
        return path

    def showBill(self):
        self.hide()
        self.bill = Bill_window()
        self.bill.show()
        AccountBook = load_workbook(login.user_file)
        sheet = AccountBook.active
        for i in range(0, sheet.max_row):
            self.bill.tableWidget.insertRow(i)
            self.check = QCheckBox()
            self.check.setChecked(False)
            self.bill.tableWidget.setCellWidget(i, 4, self.check)
            for j in range(0, 4):
                self.bill.tableWidget.setItem(i, j, QTableWidgetItem(str(sheet.cell(sheet.max_row-i, j+1).value)))

    def showStatistic(self):
        self.hide()
        self.sta = Statistic_window()
        self.sta.show()

    def undo(self):
        AccountBook = load_workbook(login.user_file)
        sheet = AccountBook.active
        sheet.delete_rows(sheet.max_row)
        AccountBook.save(login.user_file)

    def switch(self):
        self.hide()
        self.log = Login_window()
        self.log.show()


class Bill_window(QWidget, Ui_Bill):
    def __init__(self):
        super(Bill_window, self).__init__()
        self.setupUi(self)
        self.pushButton.clicked.connect(self.goback)
        self.tableWidget.setHorizontalHeaderLabels(["日期", "类型", "金额", "备注", "操作"])
        self.tableWidget.verticalHeader().setVisible(False)
        self.pushButton_2.clicked.connect(self.search)
        self.pushButton_3.clicked.connect(self.save)
        self.pushButton_4.clicked.connect(self.delete)

    def goback(self):
        self.hide()
        self.main = myMainWindow()
        self.main.show()
        self.save()

    def save(self):
        AccountBook = load_workbook(login.user_file)
        sheet = AccountBook.active
        for i in range(0, self.tableWidget.rowCount()):
            for j in range(0, 4):
                if str(sheet.cell(sheet.max_row-i, j+1).value) != self.tableWidget.item(i, j).text():
                    sheet.cell(sheet.max_row-i, j+1).value = self.tableWidget.item(i, j).text()
                    AccountBook.save(login.user_file)

    def delete(self):
        AccountBook = load_workbook(login.user_file)
        sheet = AccountBook.active
        n = self.tableWidget.rowCount()
        for i in range(n-1, -1, -1):
            if self.tableWidget.cellWidget(i, 4).isChecked():
                self.tableWidget.removeRow(i)
                sheet.delete_rows(self.tableWidget.rowCount()-i+1, 1)
        AccountBook.save(login.user_file)

    def search(self):
        self.tableWidget.setRowCount(0)
        date = self.lineEdit_2.text()
        type_ = self.lineEdit.text()
        AccountBook = load_workbook(login.user_file)
        sheet = AccountBook.active
        for i in range(1, sheet.max_row+1):
            if (date in sheet.cell(i, 1).value or date == "") \
                    and (sheet.cell(i, 2).value == type_ or type_ == ""):
                self.tableWidget.insertRow(0)
                for j in range(1, 5):
                    self.tableWidget.setItem(0, j-1, QTableWidgetItem(str(sheet.cell(i, j).value)))
                self.check = QCheckBox()
                self.check.setChecked(False)
                self.tableWidget.setCellWidget(0, 4, self.check)


class Statistic_window(QWidget, Ui_Statistic):
    def __init__(self):
        super(Statistic_window, self).__init__()
        self.setupUi(self)
        self.thisday.clicked.connect(self.today)
        self.thisweek.clicked.connect(self.week)
        self.thismonth.clicked.connect(self.month)
        self.thisyear.clicked.connect(self.year)
        self.total.clicked.connect(self.sum)
        self.pushButton_6.clicked.connect(self.goback)

    def goback(self):
        self.hide()
        self.main = myMainWindow()
        self.main.show()

    def today(self):
        date = QDate.currentDate()
        date_str = date.toString(Qt.ISODate)
        AccountBook = load_workbook(login.user_file)
        sheet = AccountBook.active
        total = 0
        x_axis = []
        y_axis = []
        for i in range(1, sheet.max_row+1):
            if sheet.cell(i, 1).value == date_str:
                total += int(sheet.cell(i, 3).value)
                x_axis.insert(0, str(sheet.max_row+1-i))
                y_axis.insert(0, int(sheet.cell(i, 3).value))
        self.textBrowser.clear()
        self.textBrowser.append(str(total))
        self.graph(x_axis, y_axis)

    def week(self):
        date = QDate.currentDate()
        date_str = date.toString(Qt.ISODate)
        AccountBook = load_workbook(login.user_file)
        sheet = AccountBook.active
        total = 0
        x_axis = []
        y_axis = [0, 0, 0, 0, 0, 0, 0]
        date1 = time.strptime(date_str, "%Y-%m-%d")
        date1 = datetime.datetime(date1[0], date1[1], date1[2])
        for i in range(1, sheet.max_row+1):
            date2 = sheet.cell(i, 1).value
            date2 = time.strptime(date2, "%Y-%m-%d")
            date2_ = datetime.datetime(date2[0], date2[1], date2[2])
            if (date1 - date2_).days < 7:
                if date2[2] not in x_axis:
                    x_axis.append(date2[2])
                total += int(sheet.cell(i, 3).value)
                y_axis[x_axis.index(date2[2])] += int(sheet.cell(i, 3).value)
        self.textBrowser.clear()
        self.textBrowser.append(str(total))
        self.graph(x_axis, y_axis)

    def month(self):
        date = QDate.currentDate()
        date_str = date.toString(Qt.ISODate)
        AccountBook = load_workbook(login.user_file)
        sheet = AccountBook.active
        total = 0
        x_axis = []
        y_axis = []
        date1 = time.strptime(date_str, "%Y-%m-%d")
        for i in range(1, sheet.max_row+1):
            date2 = sheet.cell(i, 1).value
            date2 = time.strptime(date2, "%Y-%m-%d")
            if date1[0] == date2[0] and date1[1] == date2[1]:
                if date2[2] not in x_axis:
                    x_axis.append(date2[2])
                    y_axis.append(0)
                total += int(sheet.cell(i, 3).value)
                y_axis[x_axis.index(date2[2])] += int(sheet.cell(i, 3).value)
        self.textBrowser.clear()
        self.textBrowser.append(str(total))
        self.graph(x_axis, y_axis)

    def year(self):
        date = QDate.currentDate()
        date_str = date.toString(Qt.ISODate)
        AccountBook = load_workbook(login.user_file)
        sheet = AccountBook.active
        total = 0
        x_axis = []
        y_axis = []
        date1 = time.strptime(date_str, "%Y-%m-%d")
        for i in range(1, sheet.max_row+1):
            date2 = sheet.cell(i, 1).value
            date2 = time.strptime(date2, "%Y-%m-%d")
            if date1[0] == date2[0]:
                if date2[1] not in x_axis:
                    x_axis.append(date2[1])
                    y_axis.append(0)
                total += int(sheet.cell(i, 3).value)
                y_axis[date2[1]-1] += int(sheet.cell(i, 3).value)
        self.textBrowser.clear()
        self.textBrowser.append(str(total))
        self.graph(x_axis, y_axis)

    def sum(self):
        AccountBook = load_workbook(login.user_file)
        sheet = AccountBook.active
        total = 0
        x_ = []
        x_axis = []
        y_axis = []
        for i in range(1, sheet.max_row+1):
            date2 = sheet.cell(i, 1).value
            date2 = time.strptime(date2, "%Y-%m-%d")
            total += int(sheet.cell(i, 3).value)
            if str(date2[0])+str(date2[1]) not in x_:
                x_.append(str(date2[0])+str(date2[1]))
                x_axis.append(str(date2[1]))
                y_axis.append(0)
            y_axis[len(x_axis)-1] += int(sheet.cell(i, 3).value)
        self.textBrowser.clear()
        self.textBrowser.append(str(total))
        self.graph(x_axis, y_axis)

    def graph(self, x, y):
        for i in range(len(x)):
            plt.bar(x[i], y[i], color='dodgerblue')
        for a, b, i in zip(x, y, range(len(x))):
            if y[i]:
                plt.text(a, b + b/100, "%.2f" % y[i], ha='center', fontsize=12)
        plt.xticks(x)
        plt.savefig("graph.png")
        pix = QPixmap("graph.png")
        self.label_2.setPixmap(QPixmap(""))
        self.label_2.setPixmap(pix)
        self.label_2.setScaledContents(True)
        plt.clf()


if __name__ == "__main__":
    app = QApplication(sys.argv)
    login = Login_window()
    login.show()
    sys.exit(app.exec_())
